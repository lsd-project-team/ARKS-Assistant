"""从同一正文生成两册Word、Markdown及维护索引。"""
from pathlib import Path
from importlib import import_module
from hashlib import sha256
import argparse
import json

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from figures import generate

ROOT = Path(__file__).resolve().parents[4]
BOOKS = ('tu_content', 'correction_content')


def set_font(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:'+attr), 'Microsoft YaHei')


def field(paragraph, instruction):
    for tag, value in [('begin',None),('instruction',instruction),('separate',None),('end',None)]:
        run=paragraph.add_run()
        if tag=='instruction':
            node=OxmlElement('w:instrText');node.set(qn('xml:space'),'preserve');node.text=' '+value+' '
        else:
            node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),tag)
        run._r.append(node)


def setup_document(book):
    doc=Document();section=doc.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=Mm(17);section.bottom_margin=Mm(16)
    section.left_margin=Mm(19);section.right_margin=Mm(19)
    section.header_distance=Mm(7);section.footer_distance=Mm(8)
    for name,size,bold in [('Normal',11.5,False),('Title',29,True),('Subtitle',13,False),('Heading 1',20,True),('Caption',9.5,False),('Header',9,False),('Footer',9,False)]:
        set_font(doc.styles[name],size,bold)
    for name in ['Normal','Caption','Title','Subtitle','Heading 1']:
        fmt=doc.styles[name].paragraph_format
        fmt.first_line_indent=Mm(0);fmt.left_indent=Mm(0);fmt.right_indent=Mm(0)
        fmt.space_before=Pt(0);fmt.space_after=Pt(7)
        fmt.widow_control=True
        fmt.line_spacing=Pt(18 if name=='Normal' else 15 if name=='Caption' else 35 if name=='Title' else 24)
    # python-docx默认模板的Title带有蓝色底边，需从样式中移除。
    for style in doc.styles:
        for node in list(style.element.iter(qn('w:pBdr'))):
            node.getparent().remove(node)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
    for name in ['toc 1','toc 2']:
        if name not in doc.styles:doc.styles.add_style(name,1)
        set_font(doc.styles[name],10.5)
        fmt=doc.styles[name].paragraph_format
        fmt.line_spacing=Pt(17);fmt.space_after=Pt(0);fmt.keep_with_next=False
    header=section.header.paragraphs[0]
    header.text='ZX2200    '+book.TITLE
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('第 ');field(footer,'PAGE');footer.add_run(' 页')
    settings=doc.settings.element
    node=OxmlElement('w:doNotAutoCompressPictures');settings.append(node)
    node=OxmlElement('w:updateFields');node.set(qn('w:val'),'true');settings.append(node)
    doc.core_properties.title=book.TITLE
    doc.core_properties.subject=book.INTRO
    doc.core_properties.author='ARKS'
    doc.core_properties.comments=''
    return doc


def paragraph(doc,text,bold=False,style=None):
    p=doc.add_paragraph(style=style)
    run=p.add_run(text)
    if bold:run.bold=True
    return p


def add_figure(doc,path,description,mini=False):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    # 正文使用固定行高，插图段落必须独立采用自动行高，否则Word会裁切图像。
    p.paragraph_format.line_spacing=1.0
    p.paragraph_format.space_after=Pt(3 if not mini else 7)
    p.paragraph_format.keep_with_next=True
    shape=p.add_run().add_picture(str(path),width=Mm(171))
    shape._inline.docPr.set('descr',description)


def add_cover(doc,book,out):
    p=paragraph(doc,book.TITLE,style='Title')
    p.paragraph_format.space_before=Pt(42)
    paragraph(doc,book.INTRO,style='Subtitle')
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(18)
    add_figure(doc,out/'figures'/f'{book.PAGES[0]["diagram"]}.png','阅读路线 '+book.PAGES[0]['caption'])
    paragraph(doc,'从第一页开始顺着看，每一页只讲一件事。先看设备在做什么，再看图中的记录或位置怎样改变。')
    paragraph(doc,'TU册一直跟踪“基板1、基岛2、A点”；纠偏册用独立数字例子逐步算。图中布局为示意，例子数值用于理解，不是本机参数。')


def source_report(book):
    previous=ROOT/'docs/design'/book.PREVIOUS
    original=json.loads((previous/'source-baseline.json').read_text(encoding='utf-8'))
    checked=[]
    for item in original['files']:
        actual=sha256((ROOT/item['path']).read_bytes()).hexdigest()
        if actual!=item['sha256']:raise ValueError('需要重新读受影响源码：'+item['path'])
        checked.append(dict(path=item['path'],sha256=actual))
    return dict(date='2026-10-26',revision=book.REVISION,scope='trunk/AKRS.ZX2200',
                prior_baseline=f'docs/design/{book.PREVIOUS}/source-baseline.json',
                verification='本次重写沿用旧版静态核查结论，逐文件重验哈希；不声称本次重新通读所有源码',
                files=checked,
                topics=[dict(page=i+3,title=p['title'],previous_sections=p['old']) for i,p in enumerate(book.PAGES)])


def write_sources(book,out,report):
    (out/'source-continuity.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
    previous=ROOT/'docs/design'/book.PREVIOUS
    markdown=[f'# {book.TITLE} 阅读索引','',
        '本册按设备动作重写，正文先讲场景、例子和图。详细接口与边界仍可从上版按章查阅。',
        '本次只修改文档与生成源，没有运行设备或整套C#业务；静态结论、说明模型和实机行为不混作同一项验证。',
        f'旧版核查的{len(report["files"])}个依据文件哈希均与当前源码一致。旧版本和配套源码索引保留。','',
        '| 新版主题 | 对应旧版章节 |','|---|---|']
    for i,p in enumerate(book.PAGES,1):
        markdown.append(f'| {i:02d} {p["title"]} | {p["old"]} |')
    markdown += ['',f'- [上版完整技术说明](../{previous.name}/{book.STEM}.md)',
                 f'- [上版源码入口](../{previous.name}/source-guide.md)',
                 f'- [上版源码核查发现](../{previous.name}/source-findings.md)',
                 '- [本次源码哈希与主题对照](source-continuity.json)','']
    (out/'阅读索引.md').write_text('\n'.join(markdown),encoding='utf-8',newline='\n')


def build(book):
    out=ROOT/'docs/design'/book.FOLDER
    out.mkdir(parents=True,exist_ok=True)
    report=source_report(book)
    figures=generate(book,ROOT)
    doc=setup_document(book);add_cover(doc,book,out)
    doc.add_page_break()
    paragraph(doc,'按设备流程阅读',style='Title')
    paragraph(doc,'先读完整流程，再按遇到的问题回看对应页。')
    field(doc.add_paragraph(),'TOC \\o "1-1" \\h \\z \\u')
    md=[f'# {book.TITLE}','',book.INTRO,'',
        '图中布局为示意，数字例子用于理解，不是本机参数。','']
    for number,page in enumerate(book.PAGES,1):
        doc.add_page_break()
        heading=f'{number:02d} {page["title"]}'
        paragraph(doc,heading,style='Heading 1')
        paragraph(doc,page['lead'])
        description=f'图{number} {page["caption"]}'
        add_figure(doc,out/'figures'/f'{page["diagram"]}.png',description+'。'+page['explain'])
        paragraph(doc,description,style='Caption')
        paragraph(doc,page['explain'])
        paragraph(doc,page['detail'])
        if page['extra']:
            add_figure(doc,out/'figures'/f'{page["extra"]}.png',page['remember'],True)
        paragraph(doc,page['remember'],True)
        md.extend([f'# {heading}','',page['lead'],'',f'![{description}](figures/{page["diagram"]}.png)','',description,'',page['explain'],'',page['detail'],''])
        if page['extra']:md.extend([f'![{page["remember"]}](figures/{page["extra"]}.png)',''])
        md.extend([f'**{page["remember"]}**',''])
    doc.save(out/f'{book.STEM}.docx')
    (out/f'{book.STEM}.md').write_text('\n'.join(md),encoding='utf-8',newline='\n')
    write_sources(book,out,report)
    metadata=dict(book=book.TITLE,revision=book.REVISION,topics=len(book.PAGES),
                  expected_pages=len(book.PAGES)+2,unique_figures=len(figures),
                  placements=len(figures)+1,checked_sources=len(report['files']))
    (out/'build-metadata.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
    print(json.dumps(metadata,ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--book',choices=['tu','correction','all'],default='all')
    args=parser.parse_args()
    for name in BOOKS:
        if args.book=='tu' and name!='tu_content':continue
        if args.book=='correction' and name!='correction_content':continue
        build(import_module(name))
