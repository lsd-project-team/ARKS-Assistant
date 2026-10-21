from pathlib import Path
import re,json,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from manuscript import BLOCKS
from evidence import build as evidence_build,GROUPS

TMP=Path(__file__).resolve().parents[5]/'output/design-publication-v7/correction-process';TMP.mkdir(parents=True,exist_ok=True);FIG=Path(__file__).resolve().parents[5]/'docs/design/correction-process/v3/figures'
base=evidence_build();fmap={f['id']:f for f in base['files']}
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.975);sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.73)
sec.header_distance=Inches(.32);sec.footer_distance=Inches(.3);sec.different_first_page_header_footer=True

def font(style,size,east='SimSun',latin='Times New Roman',bold=False):
    style.font.name=latin;style.font.size=Pt(size);style.font.bold=bold;style.font.italic=False;style.font.color.rgb=RGBColor(0,0,0)
    rp=style._element.get_or_add_rPr();rf=rp.get_or_add_rFonts()
    # Theme font attributes override explicit faces in Word.  Remove them so
    # reopening or updating fields cannot silently restore Calibri/MS Gothic.
    for key in list(rf.attrib):
        if 'theme' in key.lower():del rf.attrib[key]
    for key in ['ascii','hAnsi','cs']:rf.set(qn('w:'+key),latin)
    rf.set(qn('w:eastAsia'),east)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-US');lang.set(qn('w:eastAsia'),'zh-CN');rp.append(lang)
    for it in style._element.xpath('.//w:pBdr'):it.getparent().remove(it)
    pf=style.paragraph_format;pf.first_line_indent=Pt(0);pf.space_after=Pt(5);pf.widow_control=True
    snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');style._element.get_or_add_pPr().append(snap)

for name,size,east,bold in [('Normal',11.5,'SimSun',False),('Title',27,'Microsoft YaHei',True),('Subtitle',12,'Microsoft YaHei',False),('Heading 1',18.5,'Microsoft YaHei',True),('Heading 2',13,'Microsoft YaHei',True),('Heading 3',11.5,'Microsoft YaHei',True),('Caption',9.5,'SimSun',False)]:
    font(doc.styles[name],size,east,latin=east if east=='Microsoft YaHei' else 'Times New Roman',bold=bold)
normal=doc.styles['Normal'];normal.paragraph_format.line_spacing=Pt(18.5);normal.paragraph_format.first_line_indent=Pt(23);normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.paragraph_format.line_spacing=Pt(26 if name=='Heading 1' else 21);st.paragraph_format.keep_with_next=True
    st.paragraph_format.space_before=Pt(0 if name=='Heading 1' else 11);st.paragraph_format.space_after=Pt(13 if name=='Heading 1' else 6)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
font(doc.styles['Caption'],9.5);doc.styles['Caption'].paragraph_format.line_spacing=Pt(14);doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name,size in [('Book Note',9.5),('Book Source',9.0),('Book Ref Heading',11),('Book Figure Note',9.5)]:
    if name not in doc.styles:doc.styles.add_style(name,1)
    font(doc.styles[name],size,'Microsoft YaHei' if name=='Book Ref Heading' else 'SimSun',latin='Microsoft YaHei' if name=='Book Ref Heading' else 'Times New Roman',bold=name=='Book Ref Heading')
    pf=doc.styles[name].paragraph_format;pf.line_spacing=Pt(15);pf.space_after=Pt(5);pf.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Book Ref Heading'].paragraph_format.keep_with_next=True;doc.styles['Book Ref Heading'].paragraph_format.space_before=Pt(12)
if 'Book Front Heading' not in doc.styles:doc.styles.add_style('Book Front Heading',1)
font(doc.styles['Book Front Heading'],18.5,'Microsoft YaHei',latin='Microsoft YaHei',bold=True)
doc.styles['Book Front Heading'].paragraph_format.line_spacing=Pt(26)
doc.styles['Book Front Heading'].paragraph_format.space_after=Pt(16)
doc.styles['Book Front Heading'].paragraph_format.keep_with_next=True
for name,size in [('toc 1',11),('toc 2',10.5)]:
    if name not in doc.styles:doc.styles.add_style(name,1,builtin=True)
    font(doc.styles[name],size,'Microsoft YaHei',latin='Microsoft YaHei',bold=name=='toc 1')
    doc.styles[name].paragraph_format.line_spacing=Pt(18)
    doc.styles[name].paragraph_format.space_after=Pt(5 if name=='toc 1' else 1)
    doc.styles[name].paragraph_format.left_indent=Inches(.0 if name=='toc 1' else .25)
    doc.styles[name].paragraph_format.tab_stops.add_tab_stop(Inches(6.48),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)

doc.styles.add_style('Book Index',1)
font(doc.styles['Book Index'],10.5,'Microsoft YaHei',latin='Microsoft YaHei')
doc.styles['Book Index'].paragraph_format.line_spacing=Pt(18)
doc.styles['Book Index'].paragraph_format.tab_stops.add_tab_stop(Inches(6.48),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)

# Native Word outline numbering; the manuscript's readable labels stay in Markdown.
numbering=doc.part.numbering_part.element
abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),'72')
multi=OxmlElement('w:multiLevelType');multi.set(qn('w:val'),'multilevel');abstract.append(multi)
for level in range(3):
    lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),str(level))
    for tag,val in [('start','1'),('numFmt','decimal'),('lvlText','.'.join('%'+str(i+1) for i in range(level+1))),('lvlJc','left'),('suff','space')]:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);lvl.append(e)
    abstract.append(lvl)
    nr=OxmlElement('w:rPr');nf=OxmlElement('w:rFonts')
    for key in ['ascii','hAnsi','eastAsia','cs']:nf.set(qn('w:'+key),'Microsoft YaHei')
    nr.append(nf);nb=OxmlElement('w:b');nb.set(qn('w:val'),'1');nr.append(nb)
    ns=OxmlElement('w:sz');ns.set(qn('w:val'),str([37,26,23][level]));nr.append(ns);lvl.append(nr)
numbering.append(abstract);num=OxmlElement('w:num');num.set(qn('w:numId'),'72');an=OxmlElement('w:abstractNumId');an.set(qn('w:val'),'72');num.append(an);numbering.append(num)

bmcount=1000
def bookmark(p,name):
    global bmcount
    bmcount+=1
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(bmcount));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(bmcount));p._p.insert(0,a);p._p.append(b)
def link(p,label,anchor):
    h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),anchor)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'245B87');rp.append(co);r.append(rp)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def rich(p,t):
    if '**' in t:
        for piece in re.split(r'(\*\*[^*]+\*\*)',t):
            if piece.startswith('**') and piece.endswith('**'):p.add_run(piece[2:-2]).bold=True
            else:rich(p,piece)
        return
    # Preserve cited group links and keep long Latin identifiers legible.
    for s in re.split(r'(\[[1-9]\d*\])',t):
        if re.fullmatch(r'\[[1-9]\d*\]',s):link(p,s,'ref_'+s[1:-1]);continue
        for segment in re.split(r'((?<![A-Za-z0-9_])(?:p_o|p_n|U_o|U_n|a_0|a_off|T_now|T_goal|R_d)(?![A-Za-z0-9_]))',s):
            if re.fullmatch(r'(?:p_o|p_n|U_o|U_n|a_0|a_off|T_now|T_goal|R_d)',segment):
                a,b=segment.split('_');om=OxmlElement('m:oMath');ss=OxmlElement('m:sSub')
                for tag,value in [('e',a),('sub',b)]:
                    part=OxmlElement('m:'+tag);mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text=value;mr.append(mt);part.append(mr);ss.append(part)
                om.append(ss);p._p.append(om)
                continue
            for a in re.split(r'([A-Za-z_][A-Za-z0-9_.]{7,})',segment):
                r=p.add_run(a)
                if p.style.name=='Normal' and re.fullmatch(r'[A-Za-z_][A-Za-z0-9_.]{7,}',a):r.font.name='Cambria';r.font.size=Pt(10.5)
def para(t,style=None):
    p=doc.add_paragraph(style=style);rich(p,t);return p
def field(p,instr):
    r=p.add_run();beg=OxmlElement('w:fldChar');beg.set(qn('w:fldCharType'),'begin');r._r.append(beg)
    r=p.add_run();ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=' '+instr+' ';r._r.append(ins)
    r=p.add_run();sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
    r=p.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)

hp=sec.header.paragraphs[0];hp.style=doc.styles['Book Note'];hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.55),WD_TAB_ALIGNMENT.RIGHT)
hp.add_run('ZX2200 · 视觉定位与芯片纠偏\t纠偏过程设计文档')
for r in hp.runs:r.font.name='Microsoft YaHei';r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.paragraph_format.first_line_indent=Pt(0);field(fp,'PAGE')
for r in fp.runs:r.font.size=Pt(9)
doc.core_properties.title='纠偏过程设计文档';doc.core_properties.subject='ZX2200产品定位、上视与中转台芯片纠偏';doc.core_properties.author='';doc.core_properties.version='3.0'

md=['# 纠偏过程设计文档','','ZX2200视觉定位与芯片纠偏 · 第3版 · 2026年10月21日','']
p=para('ZX2200\n视觉定位与芯片纠偏','Subtitle');p.paragraph_format.space_before=Pt(58);p.paragraph_format.space_after=Pt(28)
p=para('纠偏过程','Title');p.paragraph_format.line_spacing=Pt(38);p.paragraph_format.space_after=Pt(4)
p=para('设计文档','Title');p.paragraph_format.line_spacing=Pt(38);p.paragraph_format.space_after=Pt(22)
p=para('供料定位 产品定位与芯片纠偏','Subtitle');p.paragraph_format.space_after=Pt(54)
para('解释纠偏对象、进入条件、采集步骤、坐标计算、结果使用与失败边界，让设计评审和设备调试使用同一套判断依据。')
p=para('适用读者\n软件研发与技术评审人员\n设备调试与应用工程师','Book Note');p.paragraph_format.space_before=Pt(36);p.paragraph_format.line_spacing=Pt(19)
p=para('第3版    2026年10月21日','Book Note');p.paragraph_format.space_before=Pt(36)
doc.add_page_break();para('目录','Book Front Heading');field(doc.add_paragraph(),'TOC \\o "1-2" \\h \\z \\u')

def heading(text,level,numbered=True):
    p=doc.add_paragraph(style='Heading '+str(level))
    if numbered:
        clean=re.sub(r'^\d+(?:\.\d+)*\s+','',text)
        np=OxmlElement('w:numPr');il=OxmlElement('w:ilvl');il.set(qn('w:val'),str(level-1));ni=OxmlElement('w:numId');ni.set(qn('w:val'),'72');np.append(il);np.append(ni);p._p.get_or_add_pPr().append(np)
    else:clean=text
    p.add_run(clean);bookmark(p,'h_'+str(len(doc.paragraphs)))
    return p

def caption_anchor(caption):
    code=caption.split()[0].replace('图','fig_').replace('表','tbl_').replace('-','_')
    return code

def caption_entry(p,caption,category):
    bookmark(p,caption_anchor(caption))

def picture(key,caption,description,width):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(FIG/(key+'.png')),width=Inches(width));pr=p._p.xpath('.//wp:docPr')[0];pr.set('descr',caption+'。'+description)
    q=para(caption,'Caption');q.paragraph_format.keep_with_next=True;caption_entry(q,caption,'F')
    q=para(description,'Book Figure Note');q.paragraph_format.keep_together=True;q.paragraph_format.space_after=Pt(9)
    md.extend([f'![{caption}](figures/{key}.png)','',caption,'',description,''])

def table(caption,head,rows,widths):
    q=para(caption,'Caption');q.paragraph_format.keep_with_next=True;caption_entry(q,caption,'T')
    t=doc.add_table(rows=1,cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for n,data in enumerate([head]+rows):
        row=t.rows[0] if n==0 else t.add_row();rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
        if n==0:rp.append(OxmlElement('w:tblHeader'))
        for i,text in enumerate(data):
            cell=row.cells[i];cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side,value in [('top',90),('bottom',90),('left',85),('right',85)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            if n==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EBF0F4');cp.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(15);p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=n==0;p.paragraph_format.keep_together=True
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            if caption.startswith(('表3-2','表4-1','表8-1','表9-1','表9-2')):p.paragraph_format.keep_with_next=n<len(rows)
            rich(p,text.replace('ActionNodeMessage.','ActionNodeMessage.\u200b').replace('(','(\u200b'))
            for r in p.runs:r.font.size=Pt(10);r.font.bold=n==0
    q=doc.add_paragraph();q.paragraph_format.first_line_indent=Pt(0);q.paragraph_format.space_after=Pt(2);q.paragraph_format.line_spacing=Pt(3);q.add_run().font.size=Pt(1)
    md.extend([caption,'','| '+' | '.join(head)+' |','|'+'|'.join(['---']*len(head))+'|']+['| '+' | '.join(v.replace('\n','<br>') for v in r)+' |' for r in rows]+[''])

def equation(key,parts):
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(25);p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.25),WD_TAB_ALIGNMENT.CENTER);p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run('\t');om=OxmlElement('m:oMath')
    def mr(s,plain=False):
        r=OxmlElement('m:r')
        if plain:
            rp=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');rp.append(sty);r.append(rp)
        t=OxmlElement('m:t');t.text=s;r.append(t);return r
    plain=[]
    for x in parts:
        if isinstance(x,tuple):
            _,a,b=x;ss=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(a));sub=OxmlElement('m:sub');sub.append(mr(b));ss.append(e);ss.append(sub);om.append(ss);plain.append(a+'_'+b)
        else:
            for frag in re.split('(cos|sin)',x):om.append(mr(frag,frag in ['cos','sin']))
            plain.append(x)
    p._p.append(om);p.add_run('\t('+key+')');md.extend(['$$'+''.join(plain)+'\\tag{'+key+'}$$',''])


import sys
sys.path.insert(0,str(Path(__file__).resolve().parent.parent))
from appendix import method_name,selected_anchors

for kind,value in BLOCKS:
    if kind=='chapter':
        q=heading(value,1,not value.startswith('附录'))
        q.paragraph_format.page_break_before=value.startswith('1 ') or value.startswith('附录')
        q.paragraph_format.space_before=Pt(0 if q.paragraph_format.page_break_before else 20)
        md.extend(['# '+value,''])
    elif kind in ['section','sub']:
        q=heading(value,2 if kind=='section' else 3,not re.match(r'^[A-Z]\.\d',value))
        if value.startswith('B.2 '):q.paragraph_format.page_break_before=True
        md.extend([('## ' if kind=='section' else '### ')+value,''])
    elif kind=='p':
        q=para(value)
        if value.startswith('本文以2026年'):q.alignment=WD_ALIGN_PARAGRAPH.LEFT
        md.extend([value,''])
    elif kind=='figure':picture(*value)
    elif kind=='table':table(*value)
    elif kind=='equation':equation(*value)
    elif kind in ['figure_index','table_index']:
        entries=[v for k,v in BLOCKS if k==('figure' if kind=='figure_index' else 'table')]
        for v in entries:
            title=v[1] if kind=='figure_index' else v[0]
            q=doc.add_paragraph(style='Book Index');q.paragraph_format.space_after=Pt(4)
            link(q,title,caption_anchor(title));q.add_run('\t')
            field(q,'PAGEREF '+caption_anchor(title)+' \\h')
        md.extend(['- '+(v[1] if kind=='figure_index' else v[0]) for v in entries]+[''])
    elif kind=='references':
        source_guide=['# 纠偏过程源码阅读索引','','完整路径、主题与行号锚点。文件哈希以source-baseline.json为准。','']
        primary={n:ids[:2] for n,_,_,ids in GROUPS}
        for n,title,purpose,ids in GROUPS:
            q=para(f'[{n}] {title}','Book Ref Heading');bookmark(q,'ref_'+str(n));md.extend(['## '+f'[{n}] {title}','',purpose,''])
            para(purpose,'Book Source')
            brief=[]
            for sid in primary[n]:
                f=fmap[sid]
                brief.append(Path(f['path']).name+'  '+ ' / '.join(method_name(a['symbol']) for a in selected_anchors(f,n)[:2]))
            para('；'.join(brief),'Book Source')
            source_guide.extend(['## '+f'[{n}] {title}','',purpose,''])
            for sid in ids:
                f=fmap[sid];path=f['path'].replace('trunk/AKRS.ZX2200/','ZX/').replace('trunk/AKRS.Galaxy2.CoordinateSystems/CoordinateSystems/','CS/')
                selected=selected_anchors(f,n)[:3]
                if n==13 and sid in ['S29','S30']:
                    selected=[a for a in f['anchors'] if a['symbol'] in ['baseEntity.BaseInfo.IsIdentity = true;','baseEntity.BaseInfo.IdentityCode = codeResult.CodeValue;']]
                anchors='；'.join(a['symbol'].replace('public ','').replace('private ','')+'，L'+str(a['line']) for a in selected)

                if sid in primary[n]:md.extend(['- '+Path(f['path']).name+'：'+' / '.join(method_name(a['symbol']) for a in selected_anchors(f,n)[:2])])
                source_guide.extend(['- '+f['path']+'：'+anchors])
            md.append('')
            source_guide.append('')
        (TMP/'source-guide.md').write_text('\n'.join(source_guide).rstrip()+'\n',encoding='utf-8')

settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.save(TMP/'纠偏过程设计文档.docx')
(TMP/'纠偏过程设计文档.md').write_text('\n'.join(md).rstrip()+'\n',encoding='utf-8',newline='\n')
(TMP/'book-structure.json').write_text(json.dumps(dict(chapters=[v for k,v in BLOCKS if k=='chapter'],figures=[v[1] for k,v in BLOCKS if k=='figure'],tables=[v[0] for k,v in BLOCKS if k=='table'],source_files=len(base['files']),body_paragraphs=sum(k=='p' for k,v in BLOCKS)),ensure_ascii=False,indent=2),encoding='utf-8')
print('Word manuscript and maintenance source built')
