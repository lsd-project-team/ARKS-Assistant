"""读写用户台账文件；只处理数据，不执行公式或外部链接。"""

import csv
import io
import zipfile
from pathlib import PurePath

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.modules.action_timing.ledger_schemas import MAX_FILE_BYTES, MAX_ROWS, ImportRow, name_key

HEADERS = ("分类", "动作名称", "耗时（秒）", "备注")
MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def read_rows(raw, filename):
    if len(raw) > MAX_FILE_BYTES:
        raise ValueError("文件超过5MiB，请分批导入")
    suffix = PurePath(filename).suffix.lower()
    if suffix == ".csv":
        rows = _csv_rows(raw)
    elif suffix == ".xlsx":
        rows = _xlsx_rows(raw)
    else:
        raise ValueError("请选择CSV或Excel（.xlsx）文件")
    return _validate_rows(rows, csv_escaped=suffix == ".csv")


def _csv_rows(raw):
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("gb18030")
        except UnicodeDecodeError as error:
            raise ValueError("CSV编码无法识别，请另存为UTF-8 CSV") from error
    try:
        return list(_limited(csv.reader(io.StringIO(text, newline=""), strict=True)))
    except csv.Error as error:
        raise ValueError("CSV格式错误，请使用导入模板") from error


def _xlsx_rows(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 50 * 1024 * 1024:
                raise ValueError("Excel解压后过大，请缩小文件")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.worksheets[0]
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 20:
                raise ValueError("每次最多导入5000行，表格最多20列")
            result = []
            for cells in _limited(sheet.iter_rows()):
                if any(cell.data_type in {"f", "e"} for cell in cells):
                    raise ValueError("Excel包含公式或错误单元格，请先粘贴为数值后导入")
                result.append([cell.value for cell in cells])
            return result
        finally:
            workbook.close()
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("Excel文件无法读取，请使用有效的.xlsx文件") from error


def _limited(rows):
    for index, row in enumerate(rows):
        if index > MAX_ROWS:
            raise ValueError("每次最多导入5000行，请分批导入")
        yield row


def _text(value):
    return "" if value is None else str(value)


def _validate_rows(rows, *, csv_escaped=False):
    if not rows:
        raise ValueError("文件为空，请先下载并填写模板")
    headers = [name_key(str(value or "")) for value in rows[0]]
    expected = [name_key(value) for value in HEADERS]
    if any(headers.count(value) != 1 for value in expected):
        raise ValueError("表头须包含：分类、动作名称、耗时（秒）、备注；请下载导入模板")
    if any(value and value not in expected for value in headers):
        raise ValueError("文件包含模板之外的列，请仅保留分类、动作名称、耗时（秒）、备注")
    positions = [headers.index(value) for value in expected]
    valid, errors = [], []
    for number, row in enumerate(rows[1:], 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        if any(
            value is not None and str(value).strip()
            for position, value in enumerate(row)
            if position not in positions
        ):
            errors.append({"row_number": number, "message": "存在没有表头的数据列，请按模板整理"})
            continue
        values = [row[position] if position < len(row) else None for position in positions]
        category, name, duration, notes = values
        category, name, notes = (
            restore_csv_text(_text(value)) if csv_escaped else _text(value)
            for value in (category, name, notes)
        )
        try:
            valid.append(
                ImportRow(
                    row_number=number,
                    category=category,
                    name=name,
                    duration_seconds=None
                    if duration is None or _text(duration).strip() == ""
                    else duration,
                    notes=notes,
                )
            )
        except ValidationError as error:
            labels = {
                "name": "动作名称",
                "category": "分类",
                "duration_seconds": "耗时（秒）",
                "notes": "备注",
            }
            detail = "；".join(
                f"{labels.get(item['loc'][0], item['loc'][0])}：{item['msg']}"
                for item in error.errors()
            )
            errors.append({"row_number": number, "message": detail})
    if not valid and not errors:
        raise ValueError("文件没有动作数据，请填写后再导入")
    return valid, errors


def export_file(items, format):
    rows = [[row["category"], row["name"], row["duration_seconds"], row["notes"]] for row in items]
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(HEADERS)
        for row in rows:
            writer.writerow([csv_text(value) if isinstance(value, str) else value for value in row])
        return ("\ufeff" + output.getvalue()).encode("utf-8")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "动作耗时"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for key, width in (("A", 18), ("B", 30), ("C", 18), ("D", 50)):
        sheet.column_dimensions[key].width = width
    sheet.freeze_panes = "A2"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def csv_text(value):
    # 带引号导出的文本在重导入时恢复，避免分类/名称被无意改名。
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@", "'")) else value


def restore_csv_text(value):
    if value.startswith("'") and value[1:].lstrip().startswith(("=", "+", "-", "@", "'")):
        return value[1:]
    return value
