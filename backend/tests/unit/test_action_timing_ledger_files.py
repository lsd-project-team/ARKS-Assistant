"""真实CSV/XLSX往返及非法输入；无正式数据或设备依赖。"""

import csv
import io

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.action_timing.ledger_files import HEADERS, export_file, read_rows
from app.modules.action_timing.ledger_schemas import ActionInput


def csv_file(rows, encoding="utf-8-sig", *, headers=HEADERS):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue().encode(encoding)


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_export_import_preserves_zero_null_unicode_quotes_and_formula_text(format):
    items = [
        {
            "category": "取料",
            "name": "准备,动作",
            "duration_seconds": 0,
            "notes": '首行\n备注"内容',
        },
        {"category": "未分类", "name": "待填写", "duration_seconds": None, "notes": ""},
        {"category": "=分类", "name": "+动作", "duration_seconds": 0.000001, "notes": "@备注"},
        {
            "category": "'分类",
            "name": "'=原始名称",
            "duration_seconds": 1.25,
            "notes": "  =普通文本",
        },
    ]
    raw = export_file(items, format)
    rows, errors = read_rows(raw, "动作." + format)
    assert errors == []
    assert [(row.category, row.name, row.duration_seconds, row.notes) for row in rows] == [
        (item["category"], item["name"], item["duration_seconds"], item["notes"].strip())
        for item in items
    ]
    if format == "xlsx":
        book = load_workbook(io.BytesIO(raw), data_only=False)
        assert book.active["B4"].data_type == "s" and book.active["B4"].value == "+动作"
        assert book.active["A4"].data_type == "s" and book.active["A4"].value == "=分类"
        book.close()
    else:
        assert b"'+" in raw


@pytest.mark.parametrize("duration", ["-1", "NaN", "Infinity", "错误", "true"])
def test_import_reports_bad_duration_at_actual_file_row(duration):
    rows, errors = read_rows(csv_file([["", "动作", duration, ""]]), "data.csv")
    assert rows == []
    assert errors[0]["row_number"] == 2 and "耗时" in errors[0]["message"]


@pytest.mark.parametrize("duration", [True, False, float("nan"), float("inf"), -0.01])
def test_action_input_rejects_boolean_nonfinite_and_negative_duration(duration):
    with pytest.raises(ValueError):
        ActionInput(name="动作", duration_seconds=duration)


def test_csv_gb18030_empty_duration_and_header_order():
    raw = "备注,耗时(秒),动作名称,分类\r\n备注, ,取料,取料分类\r\n".encode("gb18030")
    rows, errors = read_rows(raw, "旧表.CSV")
    assert not errors and rows[0].duration_seconds is None
    assert rows[0].name == "取料" and rows[0].notes == "备注"


@pytest.mark.parametrize(
    "filename,raw",
    [
        ("file.xls", b"invalid"),
        ("bad.xlsx", b"invalid"),
        ("empty.csv", b""),
        ("no-header.csv", "动作,毫秒\n取料,100".encode()),
        ("only-header.csv", csv_file([])),
        ("large.csv", b"x" * (5 * 1024 * 1024 + 1)),
        ("too-many.csv", csv_file([["", str(index), "", ""] for index in range(5001)])),
    ],
    ids=["extension", "bad-xlsx", "empty", "headers", "no-data", "size", "rows"],
)
def test_invalid_files_fail_without_accepting_partial_data(filename, raw):
    with pytest.raises(ValueError):
        read_rows(raw, filename)


def test_excel_formula_and_xml_control_characters_are_not_accepted():
    book = Workbook()
    book.active.append(HEADERS)
    book.active.append(["取料", "动作", "=1+1", ""])
    output = io.BytesIO()
    book.save(output)
    book.close()
    with pytest.raises(ValueError, match="公式"):
        read_rows(output.getvalue(), "formula.xlsx")
    with pytest.raises(ValueError, match="控制字符"):
        ActionInput(name="非法\x00动作")
    with pytest.raises(ValueError, match="控制字符"):
        ActionInput(name="非法\ufffe动作")


def test_excel_boolean_is_not_treated_as_one_second():
    book = Workbook()
    book.active.append(HEADERS)
    book.active.append(["", "动作", True, ""])
    output = io.BytesIO()
    book.save(output)
    book.close()
    rows, errors = read_rows(output.getvalue(), "boolean.xlsx")
    assert rows == [] and errors[0]["row_number"] == 2


def test_extra_columns_are_reported_instead_of_silently_dropping_data():
    rows, errors = read_rows(csv_file([["取料", "动作", 1, "备注", "意外多出的内容"]]), "long.csv")
    assert not rows and errors[0]["row_number"] == 2
    with pytest.raises(ValueError, match="模板之外"):
        read_rows("分类,动作名称,耗时（秒）,备注,来源\n取料,动作,1,备注,现场".encode(), "extra.csv")


@pytest.mark.parametrize("format", ["csv", "xlsx"])
@pytest.mark.parametrize(
    "header,unit,expected",
    [
        ("耗时（毫秒）", "auto", [1.25, 0, None, 0.000125]),
        ("耗时(ms)", "auto", [1.25, 0, None, 0.000125]),
        ("耗时（秒）", "auto", [1250, 0, None, 0.125]),
        ("耗时(s)", "auto", [1250, 0, None, 0.125]),
        ("耗时", "ms", [1.25, 0, None, 0.000125]),
        ("耗时", "s", [1250, 0, None, 0.125]),
    ],
)
def test_import_normalizes_explicit_units_to_seconds_before_preview(format, header, unit, expected):
    headers = ["分类", "动作名称", header, "备注"]
    values = [
        ["固晶", name, duration, ""]
        for name, duration in [("取片", 1250), ("零值", 0), ("待填", None), ("小数", 0.125)]
    ]
    if format == "csv":
        raw = csv_file(values, headers=headers)
    else:
        book = Workbook()
        book.active.append(headers)
        for row in values:
            book.active.append(row)
        output = io.BytesIO()
        book.save(output)
        book.close()
        raw = output.getvalue()
    rows, errors = read_rows(raw, "耗时." + format, unit=unit)
    assert errors == []
    assert [row.duration_seconds for row in rows] == expected
    assert [row.row_number for row in rows] == [2, 3, 4, 5]


@pytest.mark.parametrize(
    "header,unit,message",
    [
        ("耗时", "auto", "未标明单位"),
        ("耗时（秒）", "ms", "与表头不一致"),
        ("耗时（毫秒）", "s", "与表头不一致"),
        ("耗时（秒）", "minutes", "耗时单位"),
    ],
)
def test_import_rejects_ambiguous_or_conflicting_units_instead_of_guessing(header, unit, message):
    with pytest.raises(ValueError, match=message):
        read_rows(
            csv_file([["", "取片", 1000, ""]], headers=["分类", "动作名称", header, "备注"]),
            "单位.csv",
            unit=unit,
        )


def test_two_duration_columns_are_rejected_and_tiny_milliseconds_do_not_become_zero():
    with pytest.raises(ValueError, match="表头须包含"):
        read_rows(
            "分类,动作名称,耗时（秒）,耗时（毫秒）,备注\n固晶,取片,1,1000,".encode(), "重复.csv"
        )
    rows, errors = read_rows(
        csv_file(
            [["", "过小", "1e-323", ""]], headers=["分类", "动作名称", "耗时（毫秒）", "备注"]
        ),
        "小数.csv",
    )
    assert rows == [] and errors[0]["row_number"] == 2
    assert "超出可保存范围" in errors[0]["message"]
